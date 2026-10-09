"""Finite rights-clean format fixtures; executed only after the trusted preflight."""
from __future__ import annotations

import copy
import json
from pathlib import Path
import sys
import zipfile

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "skills/revayat-novel/scripts"))
sys.path.insert(0, str(ROOT / "tests"))


def refuses(error_type, operation, match):
    try:
        operation()
    except error_type as error:
        if match not in str(error):
            raise AssertionError("refusal did not name the expected boundary") from error
    else:
        raise AssertionError("invalid fixture was accepted")


def run():
    # No repository imports occur in the control parent or before its admission.
    import bookir as ir
    import bookwrite
    import e2e_evidence
    import opc
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from read_docx import read_docx

    work = Path("/scratch/fixtures")
    work.mkdir()
    cases = {}
    path = work / "unsafe.docx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("../outside.xml", "<document/>")
    refuses(opc.Damaged, lambda: opc.open_package(path), "outside")
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("word/document.xml", "<document/>")
    package = opc.open_package(path)
    try:
        assert package.xml("word/document.xml").tag == "document"
    finally:
        package.archive.close()
    cases["archive_path"] = True

    sentinel = work / "sentinel.txt"
    sentinel.write_text("not manuscript data", encoding="utf-8")
    entity = ('<!DOCTYPE document [<!ENTITY external SYSTEM "' + sentinel.as_uri() + '">]>'
              '<document>&external;</document>')
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("word/document.xml", entity)
    package = opc.open_package(path)
    try:
        refuses(opc.Damaged, lambda: package.xml("word/document.xml"), "not well-formed")
    finally:
        package.archive.close()
    cases["xml_entity"] = True

    for index, value in enumerate(("0", "-1", "bad", None)):
        document = Document()
        cell = document.add_table(rows=1, cols=1).cell(0, 0)
        cell.text = "Source"
        span = OxmlElement("w:gridSpan")
        if value is not None:
            span.set(qn("w:val"), value)
        cell._tc.get_or_add_tcPr().append(span)
        document.save(path)
        refuses(ValueError, lambda: read_docx(str(path), work / f"grid-assets-{index}"),
                "invalid DOCX table grid span")
    document = Document()
    document.add_paragraph("Unchanged source")
    document.save(path)
    assert any(block.get("text") == "Unchanged source" for block in read_docx(str(path), work / "valid-assets")["blocks"])
    cases["native_grid"] = True

    document = Document()
    table = document.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "Start"
    table.cell(1, 0).text = "Must not disappear"
    for row, value in zip(table._tbl.tr_lst, ("restart", "continue")):
        merge = OxmlElement("w:vMerge")
        merge.set(qn("w:val"), value)
        row.tc_lst[0].get_or_add_tcPr().append(merge)
    document.save(path)
    refuses(ValueError, lambda: read_docx(str(path), work / "continuation-assets"),
            "populated DOCX table vertical merge continuation")
    cases["native_continuation"] = True

    book = ir.new_book(source_path="synthetic.epub", source_format="epub", title="Control", author="Fixture")
    book["blocks"].append(ir.make_block("paragraph", 1, text="Source"))
    book_path = work / "book.json"
    ir.save_book(book, book_path)
    before = book_path.read_bytes()
    broken = copy.deepcopy(book)
    broken["blocks"][0]["text"] = "[[fn:fn9999]]"
    refuses(bookwrite.Refused, lambda: bookwrite.replace(book_path, broken, actor="sandbox", expect=bookwrite.digest(before.decode("utf-8"))),
            "nothing was written")
    assert book_path.read_bytes() == before
    assert not bookwrite.journal_path(book_path).exists()
    with bookwrite.transaction(book_path, actor="sandbox") as transaction:
        transaction.book["meta"]["title_target"] = "عنوان"
    assert ir.load_book(book_path)["meta"]["title_target"] == "عنوان"
    cases["transaction"] = True

    assert e2e_evidence.admitted(work, "book.json") == book_path
    refuses(ValueError, lambda: e2e_evidence.admitted(work, "sentinel.txt"), "Unlisted")
    book_path.unlink()
    book_path.symlink_to(sentinel)
    refuses(ValueError, lambda: e2e_evidence.admitted(work, "book.json"), "Linked")
    cases["artifact"] = True
    return cases


if __name__ == "__main__":
    print(json.dumps(run()), flush=True)
