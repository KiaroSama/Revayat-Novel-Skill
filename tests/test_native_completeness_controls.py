"""Actual package controls preserve word boundaries independently of rendering."""
import zipfile

import pytest
from docx import Document

import pagedocx


@pytest.mark.parametrize('control', ['\n', '\t', '\n\n'])
def test_native_completeness_keeps_word_boundaries(tmp_path, control):
    document = Document()
    document.add_paragraph('First' + control + 'Second')
    document.add_table(rows=1, cols=1).cell(0, 0).text = 'Third\tFourth'
    path = tmp_path / 'native.docx'
    document.save(path)
    assert pagedocx.document_text(path) == 'First' + control + 'Second Third\tFourth'
    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    parts['word/document.xml'] = parts['word/document.xml'].replace(b'<w:br/>', b'').replace(b'<w:tab/>', b'')
    with zipfile.ZipFile(path, 'w') as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    assert 'FirstSecond' in pagedocx.document_text(path)
    assert 'First' + control + 'Second' not in pagedocx.document_text(path)
