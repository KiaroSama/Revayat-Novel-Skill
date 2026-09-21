"""Real package regressions exposed by the supplied short-book render."""

from copy import deepcopy

import pytest

import bookir as ir
from build_docx import Builder
import docqa
import pagecheck
import qa
from test_package_mutations import _options


def printed(tmp_path, *, level=3, repeat=False, toc=False):
    book = ir.new_book(source_path="original.epub", source_format="epub", title="Learning a New Script")
    book["meta"]["title_target"] = "یاد گرفتن خواندن — جلد اول"
    book["blocks"] = [ir.make_block("heading", 1, level=level, text="Learning to Read",
                                    target="یاد گرفتن خواندن"),
                      ir.make_block("paragraph", 2, text="A complete paragraph.",
                                    target="این بند فارسی باید کامل و دقیقاً به تعداد تعیین‌شده در کتاب باشد.")]
    if repeat:
        repeated = deepcopy(book["blocks"][1])
        repeated["id"] = "b00003"
        book["blocks"].append(repeated)
    options = _options()
    options.toc = toc
    path = tmp_path / "book.docx"
    Builder(book, tmp_path, options).build(path)
    return book, path


def completeness(book, path):
    report = qa.Report()
    view = {"width_pt": 612, "height_pt": 792, "blocks": [], "images": []}
    docqa.check_completeness([view], docqa.document_expectations(book), report,
                             source=pagecheck.document_text(path))
    return report.summary()


@pytest.mark.parametrize("level", [3, 4, 6])
def test_headings_outside_toc_depth_still_have_bookmarks(tmp_path, level):
    book, path = printed(tmp_path, level=level)
    result = qa.check_docx(path, book).summary()
    assert result["ok"], result
    assert result["counts"]["bookmarks"] >= 1


@pytest.mark.parametrize("repeat,toc", [(False, False), (True, False), (False, True)])
def test_declared_title_and_body_repetitions_are_not_duplicates(tmp_path, repeat, toc):
    book, path = printed(tmp_path, level=1, repeat=repeat, toc=toc)
    assert completeness(book, path)["ok"]


@pytest.mark.parametrize("mutation,code", [("duplicate", "text-duplicated"), ("remove", "text-missing")])
def test_a_title_page_cannot_hide_a_missing_or_duplicated_heading(tmp_path, mutation, code):
    from docx import Document

    book, path = printed(tmp_path)
    document = Document(path)
    heading = next(p for p in document.paragraphs if p.style.name == "Heading 3")
    if mutation == "duplicate":
        heading._p.addnext(deepcopy(heading._p))
    else:
        heading._p.getparent().remove(heading._p)
    document.save(path)
    assert code in {item["code"] for item in completeness(book, path)["findings"]}


def test_a_shared_paragraph_prefix_cannot_hide_a_changed_tail(tmp_path):
    from docx import Document

    book, path = printed(tmp_path)
    common = "این بخش آغاز مشترک دو بند بلند است و باید کامل خوانده شود. " * 5
    book["blocks"][1]["target"] = common + "پایان اول."
    book["blocks"].append(ir.make_block("paragraph", 3, text="The second paragraph.",
                                        target=common + "پایان دوم."))
    Builder(book, tmp_path, _options()).build(path)
    document = Document(path)
    paragraph = next(p for p in document.paragraphs if p.text.endswith("پایان دوم."))
    paragraph.text = book["blocks"][1]["target"]
    document.save(path)
    assert "text-missing" in {item["code"] for item in completeness(book, path)["findings"]}


def test_author_subtitle_uses_the_requested_persian_face_without_theme_override(tmp_path):
    from docx import Document
    from docx.oxml.ns import qn

    book, path = printed(tmp_path)
    book["meta"].update(author="Original Author", author_target="نویسندهٔ نمونه")
    options = _options()
    Builder(book, tmp_path, options).build(path)
    fonts = Document(path).styles["Subtitle"].element.find(qn("w:rPr")).find(qn("w:rFonts"))
    assert fonts is not None and fonts.get(qn("w:cs")) == options.font
    assert fonts.get(qn("w:cstheme")) is None
