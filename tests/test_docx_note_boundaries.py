"""Native note boundaries and table ownership survive the public reader seam."""
from __future__ import annotations

import io

import pytest
from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from read_docx import read_docx
from test_docx_event_order import imported, note_part, reference
from tests_support import png_bytes


@pytest.mark.parametrize("kind", ["footnote", "endnote"])
@pytest.mark.parametrize("paragraphs,expected", [
    (["", "First", "", "Last", ""], "\nFirst\n\nLast\n"),
    ([" ", "First", " "], " \nFirst\n "),
    ([" First ", "", " Last "], " First \n\n Last "),
])
def test_note_keeps_every_authored_paragraph_and_its_edges(tmp_path, kind, paragraphs, expected):
    document = Document()
    reference(document.add_paragraph("Body").add_run(), kind)
    body = "".join('<w:p><w:r><w:t xml:space="preserve">' + text +
                   '</w:t></w:r></w:p>' for text in paragraphs)
    note_part(document, kind, [("1", body)])
    assert imported(document, tmp_path)["footnotes"][0]["text"] == expected


@pytest.mark.parametrize("kind", ["footnote", "endnote"])
def test_note_with_only_empty_paragraphs_still_refuses(tmp_path, kind):
    document = Document()
    reference(document.add_paragraph("Body").add_run(), kind)
    note_part(document, kind, [("1", "<w:p/><w:p/><w:p/>")])
    source = tmp_path / "empty-note.docx"
    document.save(source)
    with pytest.raises(ValueError, match="empty DOCX note"):
        read_docx(str(source), tmp_path / "assets")


@pytest.mark.parametrize("content", ["text", "tab", "line", "page", "image", "note", "table", "field"])
def test_populated_vertical_merge_continuation_cannot_disappear(tmp_path, content):
    document = Document()
    table = document.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "Merged source"
    cell = table.cell(1, 0)
    run = cell.paragraphs[0].add_run()
    if content == "text":
        run.add_text("Continuation source")
    elif content == "tab":
        run.add_tab()
    elif content in {"line", "page"}:
        run.add_break(WD_BREAK.LINE if content == "line" else WD_BREAK.PAGE)
    elif content == "image":
        run.add_picture(io.BytesIO(png_bytes(20, 10)))
    elif content == "note":
        reference(run)
        note_part(document)
    elif content == "table":
        cell.add_table(rows=1, cols=1).cell(0, 0).text = "Nested source"
    else:
        field = OxmlElement("w:fldChar")
        field.set(qn("w:fldCharType"), "begin")
        run._r.append(field)
    for row, value in zip(table._tbl.tr_lst, ["restart", "continue"]):
        merge = OxmlElement("w:vMerge")
        merge.set(qn("w:val"), value)
        row.tc_lst[0].get_or_add_tcPr().append(merge)
    source = tmp_path / "populated-continuation.docx"
    document.save(source)
    with pytest.raises(ValueError, match="populated DOCX table vertical merge continuation"):
        read_docx(str(source), tmp_path / "assets")


def test_generated_empty_vertical_merge_continuation_remains_supported(tmp_path):
    document = Document()
    table = document.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "Merged source"
    table.cell(0, 0).merge(table.cell(1, 0))
    book = imported(document, tmp_path)
    assert [(block["text"], block["row_span"]) for block in book["blocks"]] == [("Merged source", 2)]


@pytest.mark.parametrize("value", ["0", "-1", "bad", None])
def test_invalid_explicit_grid_span_refuses_instead_of_becoming_one(tmp_path, value):
    document = Document()
    cell = document.add_table(rows=1, cols=1).cell(0, 0)
    cell.text = "Source"
    span = OxmlElement("w:gridSpan")
    if value is not None:
        span.set(qn("w:val"), value)
    cell._tc.get_or_add_tcPr().append(span)
    source = tmp_path / "invalid-span.docx"
    document.save(source)
    with pytest.raises(ValueError, match="invalid DOCX table grid span"):
        read_docx(str(source), tmp_path / "assets")


def test_nested_events_carry_their_immediate_parent_cell(tmp_path):
    document = Document()
    parent = document.add_table(rows=1, cols=1).cell(0, 0)
    parent.paragraphs[0].text = "Parent before"
    child = parent.add_table(rows=1, cols=1).cell(0, 0)
    run = child.paragraphs[0].add_run("Child before")
    run.add_picture(io.BytesIO(png_bytes(20, 10)))
    run.add_text("Child middle")
    reference(run)
    run.add_break(WD_BREAK.PAGE)
    run.add_text("Child after")
    child.add_table(rows=1, cols=1).cell(0, 0).text = "Grandchild"
    parent.add_paragraph("Parent after")
    note_part(document)
    book = imported(document, tmp_path)
    blocks = book["blocks"]
    assert [(block["type"], block.get("text") or "") for block in blocks] == [
        ("paragraph", "Parent before"), ("paragraph", "Child before"),
        ("image", ""), ("paragraph", "Child middle[[fn:fn0001]]"),
        ("pagebreak", ""), ("paragraph", "Child after"),
        ("paragraph", "Grandchild"), ("paragraph", "Parent after"),
    ]
    assert [(block.get("table"), block.get("row"), block.get("cell"),
             block.get("parent_table"), block.get("parent_row"), block.get("parent_cell"))
            for block in blocks] == [
        ("t0000", 1, 1, None, None, None),
        *[("t0000-11n1", 1, 1, "t0000", 1, 1)] * 5,
        ("t0000-11n1-11n1", 1, 1, "t0000-11n1", 1, 1),
        ("t0000", 1, 1, None, None, None),
    ]
    assert book["footnotes"][0]["anchor_block"] == blocks[3]["id"]
