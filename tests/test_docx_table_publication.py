"""Ordered native table events must survive actual Word publication."""
from __future__ import annotations

import argparse
import io

from docx import Document
from docx.enum.text import WD_BREAK
import zipfile
from xml.etree import ElementTree as ET

import pytest

import bookir as ir
import opc
import qa
from build_docx import Builder, add_arguments
from tests_support import png_bytes
from read_docx import read_docx
import chunk
import merge
import pagerun
import worksheet


def build_table(book, work):
    parser = argparse.ArgumentParser()
    add_arguments(parser)
    opts = parser.parse_args(['--book', 'unused', '--out', 'unused', '--no-toc', '--no-justify', '--page-breaks', 'source'])
    out = work/'published.docx'
    Builder(book, work/'assets' if (work/'assets').is_dir() else work, opts).build(out)
    with zipfile.ZipFile(out) as archive:
        root = ET.fromstring(archive.read('word/document.xml'))
    return out, root


@pytest.mark.parametrize('event', ['paragraph', 'image', 'pagebreak'])
def test_cell_fragments_append_in_source_order_without_losing_events(tmp_path, event):
    book = ir.new_book(source_format='docx')
    common = {'table':'tbl0001', 'row':1, 'cell':1}
    book['blocks'] = [ir.make_block('paragraph', 1, text='Before', target='Before', **common)]
    if event == 'paragraph':
        middle = ir.make_block('paragraph', 2, text='Middle', target='Middle', **common)
    elif event == 'pagebreak':
        middle = ir.make_block('pagebreak', 2, soft=False, **common)
    else:
        data = png_bytes(20, 10)
        (tmp_path/'picture.png').write_bytes(data)
        middle = ir.make_block('image', 2, asset='picture.png', sha256=ir.sha256_bytes(data),
                               width_pt=20, height_pt=10, **common)
    book['blocks'] += [middle, ir.make_block('paragraph', 3, text='After', target='After', **common)]
    out, root = build_table(book, tmp_path)
    tables = root.find(opc.qname('w','body')).findall(opc.qname('w','tbl'))
    assert len(tables) == 1
    cell = next(tables[0].iter(opc.qname('w','tc')))
    text = opc.text_of(cell)
    assert text.startswith('Before') and text.endswith('After'), text
    if event == 'paragraph':
        paragraphs = cell.findall(opc.qname('w','p'))
        assert [opc.text_of(p) for p in paragraphs] == ['Before', 'Middle', 'After']
    elif event == 'image':
        assert len(list(cell.iter(opc.qname('w','drawing')))) == 1
    else:
        assert [n.get(opc.qname('w','type')) for n in cell.iter(opc.qname('w','br'))] == ['page']
    assert qa.check_docx(out, book).summary()['ok']


def test_nested_table_stays_inside_its_cell_between_paragraphs(tmp_path):
    book = ir.new_book(source_format='docx')
    parent = {'table':'tbl0001', 'row':1, 'cell':1}
    child = {'table':'tbl0002', 'row':1, 'cell':1,
             'parent_table':'tbl0001', 'parent_row':1, 'parent_cell':1}
    book['blocks'] = [ir.make_block('paragraph', 1, text='Before', target='Before', **parent),
                      ir.make_block('paragraph', 2, text='Inside', target='Inside', **child),
                      ir.make_block('paragraph', 3, text='After', target='After', **parent),
                      ir.make_block('paragraph', 4, text='Next cell', target='Next cell',
                                    table='tbl0001', row=1, cell=2)]
    out, root = build_table(book, tmp_path)
    tables = root.find(opc.qname('w','body')).findall(opc.qname('w','tbl'))
    assert len(tables) == 1
    cells = tables[0].find(opc.qname('w','tr')).findall(opc.qname('w','tc'))
    events = [(node.tag.rsplit('}', 1)[-1], opc.text_of(node)) for node in cells[0]
              if node.tag in {opc.qname('w','p'), opc.qname('w','tbl')}]
    assert events == [('p','Before'), ('tbl','Inside'), ('p','After')]
    assert opc.text_of(cells[1]) == 'Next cell'
    assert qa.check_docx(out, book).summary()['ok']


@pytest.mark.parametrize('route', ['chunks', 'pages'])
def test_native_nested_cell_events_survive_bound_merge_and_publication(tmp_path, route):
    source = Document()
    cell = source.add_table(rows=1, cols=2).cell(0, 0)
    cell.paragraphs[0].text = 'Before'
    picture = png_bytes(20, 10)
    run = cell.paragraphs[0].add_run()
    run.add_picture(io.BytesIO(picture))
    run.add_text('After picture')
    cell.add_table(rows=1, cols=1).cell(0, 0).text = 'Inside'
    cell.add_paragraph('After nested')
    run = cell.add_paragraph('Before break').add_run()
    run.add_break(WD_BREAK.PAGE)
    run.add_text('After break')
    source.tables[0].cell(0, 1).text = 'Next cell'
    path = tmp_path/'source.docx'
    source.save(path)
    book = read_docx(str(path), tmp_path/'assets')
    for block in book['blocks']:
        block['page'] = 1
    book['meta'].update(title='', author='')
    book['pages'] = 1
    book_path, jobs = tmp_path/'book.json', tmp_path/'jobs'
    ir.save_book(book, book_path)
    manifest = (chunk if route == 'chunks' else pagerun).build(book_path, jobs, glossary_path=None, budget=5000)
    for entry in manifest['chunks']:
        units = worksheet.read_worksheet((jobs/entry['file']).read_text(encoding='utf-8'))
        lines = [worksheet.request_line(entry['request'])]
        for unit in units:
            lines += [f"@@ {unit['id']} {unit['kind']}", worksheet.escape_payload(unit['text'])]
        ir.write_text(jobs/entry['output'], '\n'.join(lines)+'\n')
    result = merge.merge(book_path, jobs) if route == 'chunks' else pagerun.merge_page(book_path, jobs, 1)
    assert result['ok'], result
    translated = ir.load_book(book_path)
    out, root = build_table(translated, tmp_path)
    body = root.find(opc.qname('w','body'))
    assert len(body.findall(opc.qname('w','tbl'))) == 1
    cells = body.find(opc.qname('w','tbl')).find(opc.qname('w','tr')).findall(opc.qname('w','tc'))
    events = []
    for node in cells[0]:
        if node.tag == opc.qname('w','tbl'):
            events.append(('table', opc.text_of(node)))
        elif node.tag == opc.qname('w','p'):
            if list(node.iter(opc.qname('w','drawing'))):
                events.append(('image', ''))
            elif any(item.get(opc.qname('w','type')) == 'page' for item in node.iter(opc.qname('w','br'))):
                events.append(('pagebreak', ''))
            elif opc.text_of(node):
                events.append(('text', opc.text_of(node)))
    assert events == [('text','Before'), ('image',''), ('text','After picture'),
                      ('table','Inside'), ('text','After nested'), ('text','Before break'),
                      ('pagebreak',''), ('text','After break')]
    assert opc.text_of(cells[1]) == 'Next cell'
    asset = next(block['asset'] for block in translated['blocks'] if block['type']=='image')
    assert (tmp_path/'assets'/asset).read_bytes() == picture
    assert qa.check_docx(out, translated).summary()['ok']
