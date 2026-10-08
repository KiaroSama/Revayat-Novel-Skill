"""Authored native edges are content, not extraction noise (no renderer needed)."""
from docx import Document
from docx.oxml import OxmlElement

import bookir as ir
from read_docx import read_docx


def test_native_nonblank_edges_and_running_controls_survive_intake(tmp_path):
    document = Document()
    document.add_paragraph('  First\tSecond\nLast  ')
    head = document.sections[0].header.paragraphs[0]
    run = head.add_run('  Header')
    for tag, following in [('w:cr', 'Next'), ('w:noBreakHyphen', 'joined'),
                           ('w:softHyphen', 'soft')]:
        run._r.append(OxmlElement(tag))
        run.add_text(following)
    run.add_break()
    run.add_text('Last  ')
    source = tmp_path / 'native.docx'
    document.save(source)
    book = read_docx(str(source), tmp_path / 'assets')
    assert [b['text'] for b in book['blocks'] if b['type'] in ir.TEXT_TYPES] == [
        '  First\tSecond\nLast  ']
    pieces = book['sections'][0]['headers']['default']['paragraphs'][0]['pieces']
    assert [p['text'] for p in pieces if p.get('id')] == [
        '  Header\nNext‑joined­soft\nLast  ']
    assert ir.validate_book(book) == []


def test_native_affine_image_rotation_is_not_lost(tmp_path):
    import argparse
    import zipfile
    from xml.etree import ElementTree as ET
    import opc
    from build_docx import Builder, add_arguments
    from tests_support import png_bytes

    data = png_bytes(20, 10)
    (tmp_path / 'picture.png').write_bytes(data)
    book = ir.new_book()
    book['blocks'] = [ir.make_block('image', 1, asset='picture.png', sha256=ir.sha256_bytes(data),
                                   width_pt=10, height_pt=20, transform=[0, 20, -10, 0, 10, 0])]
    parser = argparse.ArgumentParser()
    add_arguments(parser)
    options = parser.parse_args(['--book', 'unused', '--out', 'unused', '--no-toc', '--no-page-numbers'])
    out = tmp_path / 'rotated.docx'
    Builder(book, tmp_path, options).build(out)
    with zipfile.ZipFile(out) as archive:
        root = ET.fromstring(archive.read('word/document.xml'))
    transform = next(root.iter(opc.qname('a', 'xfrm')))
    assert transform.get('rot') == '5400000'
    assert transform.find(opc.qname('a', 'ext')).attrib == {'cx': '254000', 'cy': '127000'}
    inline = next(root.iter(opc.qname('wp', 'inline')))
    assert inline.find(opc.qname('wp', 'extent')).attrib == {'cx': '254000', 'cy': '127000'}
    effect = inline.find(opc.qname('wp', 'effectExtent'))
    assert effect.attrib == {'l': '0', 'r': '0', 't': '63500', 'b': '63500'}
    import qa
    assert qa.check_docx(out, book).summary()['ok']
    with zipfile.ZipFile(out) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    inline.remove(effect)
    parts['word/document.xml'] = ET.tostring(root, encoding='utf-8')
    with zipfile.ZipFile(out, 'w') as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    assert 'image-transform-mismatch' in {f['code'] for f in qa.check_docx(out, book).findings}


def test_native_affine_reflection_and_mutated_orientation_are_checked(tmp_path):
    import argparse
    import zipfile
    from xml.etree import ElementTree as ET
    import opc
    import qa
    from build_docx import Builder, add_arguments
    from tests_support import png_bytes

    data = png_bytes(20, 10)
    (tmp_path / 'picture.png').write_bytes(data)
    book = ir.new_book()
    book['blocks'] = [ir.make_block('image', 1, asset='picture.png', sha256=ir.sha256_bytes(data),
                                   width_pt=20, height_pt=10, transform=[20, 0, 0, -10, 0, 10])]
    parser = argparse.ArgumentParser()
    add_arguments(parser)
    options = parser.parse_args(['--book', 'unused', '--out', 'unused', '--no-toc', '--no-page-numbers'])
    out = tmp_path / 'reflected.docx'
    Builder(book, tmp_path, options).build(out)
    assert qa.check_docx(out, book).summary()['ok']
    with zipfile.ZipFile(out) as archive:
        parts = {n: archive.read(n) for n in archive.namelist()}
    root = ET.fromstring(parts['word/document.xml'])
    transform = next(root.iter(opc.qname('a', 'xfrm')))
    assert transform.get('flipV') == '1'
    transform.set('flipV', '0')
    parts['word/document.xml'] = ET.tostring(root, encoding='utf-8')
    with zipfile.ZipFile(out, 'w') as archive:
        for name, raw in parts.items():
            archive.writestr(name, raw)
    assert 'image-transform-mismatch' in {f['code'] for f in qa.check_docx(out, book).findings}


def test_native_affine_shear_refuses_without_raster_rewrite(tmp_path):
    import argparse
    import pytest
    from build_docx import Builder, add_arguments
    from tests_support import png_bytes

    data = png_bytes(20, 10)
    (tmp_path / 'picture.png').write_bytes(data)
    book = ir.new_book()
    book['blocks'] = [ir.make_block('image', 1, asset='picture.png', sha256=ir.sha256_bytes(data),
                                   width_pt=25, height_pt=10, transform=[20, 0, 5, 10, 0, 0])]
    parser = argparse.ArgumentParser()
    add_arguments(parser)
    options = parser.parse_args(['--book', 'unused', '--out', 'unused', '--no-toc'])
    with pytest.raises(ValueError, match='unsupported image affine transform: shear'):
        Builder(book, tmp_path, options).build(tmp_path / 'refused.docx')
    assert (tmp_path / 'picture.png').read_bytes() == data
    assert not (tmp_path / 'refused.docx').exists()


def test_native_nonblank_spaces_between_drawings_are_not_discarded(tmp_path):
    import io
    from tests_support import png_bytes

    document = Document()
    run = document.add_paragraph().add_run('Before ')
    run.add_picture(io.BytesIO(png_bytes(10, 10)))
    run.add_text(' After ')
    source = tmp_path / 'inline.docx'
    document.save(source)
    book = read_docx(str(source), tmp_path / 'assets')
    assert [(b['type'], b.get('text')) for b in book['blocks']] == [
        ('paragraph', 'Before '), ('image', None), ('paragraph', ' After ')]
