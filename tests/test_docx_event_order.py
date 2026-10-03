"""Source OOXML event order must survive import without silent omissions."""
from __future__ import annotations

import io
import logging

import pytest
from docx import Document
from docx.enum.text import WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

import bookir as ir
from read_docx import read_docx
from tests_support import png_bytes

LOG = logging.getLogger(__name__)
W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'


def note_part(document, kind='footnote', bodies=None, raw=None):
    bodies = bodies if bodies is not None else [('1', '<w:p><w:r><w:t>Note body.</w:t></w:r></w:p>')]
    text = ('<w:'+kind+'s xmlns:w="'+W+'">' + ''.join(
        '<w:'+kind+' w:id="'+key+'">'+body+'</w:'+kind+'>' for key, body in bodies) + '</w:'+kind+'s>')
    part = Part(PackURI('/word/'+kind+'s.xml'),
                'application/vnd.openxmlformats-officedocument.wordprocessingml.'+kind+'s+xml',
                raw if raw is not None else text.encode(), document.part.package)
    document.part.relate_to(part, RT.FOOTNOTES if kind == 'footnote' else RT.ENDNOTES)


def reference(run, kind='footnote', identity='1'):
    element = OxmlElement('w:'+kind+'Reference')
    element.set(qn('w:id'), identity)
    run._r.append(element)


def imported(document, tmp_path):
    source = tmp_path / 'source.docx'
    document.save(source)
    book = read_docx(str(source), tmp_path / 'assets')
    assert ir.validate_book(book) == []
    LOG.debug('Imported source fixture with %d blocks', len(book['blocks']))
    return book


@pytest.mark.parametrize('kind', ['footnote', 'endnote'])
@pytest.mark.parametrize('same_run', [False, True])
def test_note_reference_remains_between_its_source_words(tmp_path, kind, same_run):
    document = Document()
    p = document.add_paragraph()
    run = p.add_run('Before ')
    reference(run if same_run else p.add_run(), kind)
    (run if same_run else p.add_run()).add_text(' after.')
    note_part(document, kind)
    book = imported(document, tmp_path)
    assert book['blocks'][0]['text'] == 'Before [[fn:fn0001]] after.'
    assert book['footnotes'][0]['anchor_block'] == book['blocks'][0]['id']


@pytest.mark.parametrize('kind', ['footnote', 'endnote'])
def test_reference_only_paragraph_is_content(tmp_path, kind):
    document = Document()
    reference(document.add_paragraph().add_run(), kind)
    note_part(document, kind)
    book = imported(document, tmp_path)
    assert len(book['blocks']) == len(book['footnotes']) == 1
    assert book['blocks'][0]['text'] == '[[fn:fn0001]]'


@pytest.mark.parametrize('kind', ['footnote', 'endnote'])
def test_each_repeated_reference_has_an_owned_note(tmp_path, kind):
    document = Document()
    for text in ['First', 'Second']:
        reference(document.add_paragraph(text).add_run(), kind)
    note_part(document, kind)
    book = imported(document, tmp_path)
    assert len(book['footnotes']) == 2
    assert [n['text'] for n in book['footnotes']] == ['Note body.'] * 2
    assert [ir.footnote_refs(b['text']) for b in book['blocks']] == [['fn0001'], ['fn0002']]


@pytest.mark.parametrize('kind', ['footnote', 'endnote'])
@pytest.mark.parametrize('defect', ['missing-part', 'missing-id', 'empty', 'malformed', 'duplicate-id'])
def test_unresolved_note_refuses_instead_of_shortening_success(tmp_path, kind, defect):
    document = Document()
    reference(document.add_paragraph('Keep the source.').add_run(), kind)
    if defect == 'missing-id':
        note_part(document, kind, [('2', '<w:p><w:r><w:t>Other note.</w:t></w:r></w:p>')])
    elif defect == 'empty':
        note_part(document, kind, [('1', '<w:p/>')])
    elif defect == 'malformed':
        note_part(document, kind, raw=b'<invalid')
    elif defect == 'duplicate-id':
        note_part(document, kind, [('1', '<w:p><w:r><w:t>A</w:t></w:r></w:p>'),
                                   ('1', '<w:p><w:r><w:t>B</w:t></w:r></w:p>')])
    source = tmp_path / 'source.docx'
    document.save(source)
    with pytest.raises(ValueError, match='note'):
        read_docx(str(source), tmp_path / 'assets')


@pytest.mark.parametrize('kind', ['footnote', 'endnote'])
def test_note_paragraphs_controls_and_style_are_not_flattened(tmp_path, kind):
    document = Document()
    reference(document.add_paragraph('Text.').add_run(), kind)
    body = ('<w:p><w:r><w:t>First</w:t><w:tab/><w:t>column</w:t><w:br/><w:t>line</w:t></w:r></w:p>'
            '<w:p><w:r><w:rPr><w:b/></w:rPr><w:t>Second</w:t></w:r></w:p>')
    note_part(document, kind, [('1', body)])
    book = imported(document, tmp_path)
    assert book['footnotes'][0]['text'] == 'First\tcolumn\nline\n**Second**'


@pytest.mark.parametrize('count', [1, 2, 3])
def test_text_and_every_drawing_in_one_run_keep_xml_order(tmp_path, count):
    document = Document()
    run = document.add_paragraph().add_run('Start')
    images = []
    for i in range(count):
        data = png_bytes(20+i, 10+i)
        images.append(data)
        run.add_picture(io.BytesIO(data), width=Pt(20+i))
        run.add_text(f'After {i}')
    book = imported(document, tmp_path)
    assert [b['type'] for b in book['blocks']] == ['paragraph'] + ['image', 'paragraph'] * count
    assert [b['text'] for b in book['blocks'] if b['type'] == 'paragraph'] == ['Start'] + [f'After {i}' for i in range(count)]
    assert [b['sha256'] for b in book['blocks'] if b['type'] == 'image'] == [ir.sha256_bytes(data) for data in images]


@pytest.mark.parametrize('count', [1, 2])
def test_hard_page_break_is_not_moved_after_later_text(tmp_path, count):
    document = Document()
    run = document.add_paragraph().add_run('Start')
    for i in range(count):
        run.add_break(WD_BREAK.PAGE)
        run.add_text(f'Page {i+2}')
    book = imported(document, tmp_path)
    assert [b['type'] for b in book['blocks']] == ['paragraph'] + ['pagebreak', 'paragraph'] * count
    assert [b['text'] for b in book['blocks'] if b['type'] == 'paragraph'] == ['Start'] + [f'Page {i+2}' for i in range(count)]


def test_failed_reimport_does_not_replace_a_previous_book_image(tmp_path):
    first = Document()
    old = png_bytes(24, 12)
    first.add_paragraph('Original').add_run().add_picture(io.BytesIO(old))
    book = imported(first, tmp_path)
    old_asset = tmp_path / 'assets' / next(b['asset'] for b in book['blocks'] if b['type'] == 'image')
    second = Document()
    run = second.add_paragraph().add_run('New')
    run.add_picture(io.BytesIO(png_bytes(42, 18)))
    reference(run, identity='999')
    source = tmp_path / 'broken.docx'
    second.save(source)
    with pytest.raises(ValueError, match='note'):
        read_docx(str(source), tmp_path / 'assets')
    assert old_asset.read_bytes() == old


def test_nested_table_stays_between_its_surrounding_cell_paragraphs(tmp_path):
    document = Document()
    cell = document.add_table(rows=1, cols=1).cell(0, 0)
    cell.paragraphs[0].text = 'Before'
    cell.add_table(rows=1, cols=1).cell(0, 0).text = 'Inside'
    cell.add_paragraph('After')
    book = imported(document, tmp_path)
    assert [b['text'] for b in book['blocks']] == ['Before', 'Inside', 'After']


@pytest.mark.parametrize('kind', ['footnote', 'endnote'])
def test_reference_after_image_in_same_run_keeps_its_position(tmp_path, kind):
    document = Document()
    run = document.add_paragraph().add_run('Before')
    run.add_picture(io.BytesIO(png_bytes(20, 10)))
    run.add_text('Next')
    reference(run, kind)
    run.add_text('Later')
    note_part(document, kind)
    book = imported(document, tmp_path)
    assert [b['type'] for b in book['blocks']] == ['paragraph', 'image', 'paragraph']
    assert book['blocks'][2]['text'] == 'Next[[fn:fn0001]]Later'
    assert book['footnotes'][0]['anchor_block'] == book['blocks'][2]['id']


@pytest.mark.parametrize('text', ['\nFirst\n\nLast\n', 'First\t\tLast', '\tFirst\t'])
def test_native_body_layout_controls_survive_exactly(tmp_path, text):
    document = Document()
    document.add_paragraph(text)
    assert imported(document, tmp_path)['blocks'][0]['text'] == text


@pytest.mark.parametrize('defect', ['missing-relationship', 'invalid-width', 'zero-width', 'missing-extent'])
def test_invalid_drawing_refuses_instead_of_disappearing(tmp_path, defect):
    document = Document()
    run = document.add_paragraph().add_run('Before')
    run.add_picture(io.BytesIO(png_bytes(20, 10)))
    drawing = run._r.find(qn('w:drawing'))
    extent = drawing.find('.//' + qn('wp:extent'))
    if defect == 'missing-relationship':
        drawing.find('.//' + qn('a:blip')).set(qn('r:embed'), 'unknown')
    elif defect == 'missing-extent':
        extent.getparent().remove(extent)
    else:
        extent.set('cx', 'bad' if defect == 'invalid-width' else '0')
    source = tmp_path / 'broken.docx'
    document.save(source)
    with pytest.raises(ValueError, match='picture'):
        read_docx(str(source), tmp_path / 'assets')


@pytest.mark.parametrize('tag', ['w:sym', 'w:object', 'w:pict'])
def test_unsupported_run_content_is_not_silently_removed(tmp_path, tag):
    document = Document()
    run = document.add_paragraph().add_run('Before')
    run._r.append(OxmlElement(tag))
    run.add_text('After')
    source = tmp_path / 'unsupported.docx'
    document.save(source)
    with pytest.raises(ValueError, match='unsupported DOCX'):
        read_docx(str(source), tmp_path / 'assets')


@pytest.mark.parametrize('tag', ['w:sdt', 'w:ins', 'm:oMath', 'w:altChunk'])
def test_unsupported_source_container_refuses_before_assets(tmp_path, tag):
    document = Document()
    paragraph = document.add_paragraph('Before')
    node = OxmlElement(tag)
    paragraph._p.append(node)
    document.add_paragraph('After')
    source = tmp_path / 'unsupported.docx'
    document.save(source)
    with pytest.raises(ValueError, match='unsupported DOCX'):
        read_docx(str(source), tmp_path / 'assets')
    assert not list((tmp_path / 'assets').glob('*'))


def test_native_cli_refusal_keeps_prior_book_and_assets(tmp_path):
    import json
    import sys
    from pathlib import Path
    import read_docx as reader

    work = tmp_path / 'work'
    work.mkdir()
    existing = ir.new_book(source_format='docx')
    existing['blocks'] = [ir.make_block('paragraph', 1, text='Existing source', target='ترجمهٔ محفوظ')]
    path = work / 'book.json'
    ir.save_book(existing, path)
    before = path.read_bytes()
    broken = Document()
    reference(broken.add_paragraph('Missing note').add_run(), identity='999')
    source = tmp_path / 'broken.docx'
    broken.save(source)
    cli = Path(reader.__file__).with_name('revayat-novel.py')
    result = ir.run_bounded([sys.executable, str(cli), 'extract', str(source), '--out', str(work)], 20)
    assert result.returncode == 2
    assert json.loads(result.stdout.decode('utf-8'))['ok'] is False
    assert 'Traceback' not in result.stderr.decode('utf-8')
    assert path.read_bytes() == before


@pytest.mark.parametrize('kind', ['footnote', 'endnote'])
def test_normal_note_zero_is_not_mistaken_for_a_separator(tmp_path, kind):
    document = Document()
    reference(document.add_paragraph('A').add_run(), kind, '0')
    note_part(document, kind, [('0', '<w:p><w:r><w:t>Normal zero.</w:t></w:r></w:p>')])
    book = imported(document, tmp_path)
    assert book['footnotes'][0]['text'] == 'Normal zero.'


def test_footnote_and_endnote_ids_are_independent(tmp_path):
    document = Document()
    run = document.add_paragraph().add_run('A')
    reference(run, 'footnote')
    run.add_text('B')
    reference(run, 'endnote')
    run.add_text('C')
    note_part(document, 'footnote')
    note_part(document, 'endnote', [('1', '<w:p><w:r><w:t>End body.</w:t></w:r></w:p>')])
    book = imported(document, tmp_path)
    assert book['blocks'][0]['text'] == 'A[[fn:fn0001]]B[[fn:fn0002]]C'
    assert [n['text'] for n in book['footnotes']] == ['Note body.', 'End body.']


def test_link_metadata_follows_only_its_actual_fragments(tmp_path):
    document = Document()
    paragraph = document.add_paragraph('Unlinked')
    paragraph.add_run().add_picture(io.BytesIO(png_bytes(10, 10)))
    hyperlink = OxmlElement('w:hyperlink')
    rid = document.part.relate_to('https://example.org/target', RT.HYPERLINK, is_external=True)
    hyperlink.set(qn('r:id'), rid)
    run = paragraph.add_run('Before')
    run.add_picture(io.BytesIO(png_bytes(20, 10)))
    run.add_text('After')
    hyperlink.append(run._r)
    paragraph._p.append(hyperlink)
    book = imported(document, tmp_path)
    prose = [b for b in book['blocks'] if b['type'] == 'paragraph']
    assert [b['text'] for b in prose] == ['Unlinked', 'Before', 'After']
    assert [b.get('links', []) for b in prose] == [[],
        [{'text': 'Before', 'href': 'https://example.org/target'}],
        [{'text': 'After', 'href': 'https://example.org/target'}]]


def test_grid_before_keeps_a_cells_real_column(tmp_path):
    document = Document()
    table = document.add_table(rows=1, cols=3)
    table.cell(0, 1).text = 'Column two'
    table.cell(0, 2).text = 'Column three'
    row = table._tbl.tr_lst[0]
    row.remove(row.tc_lst[0])
    props = row.get_or_add_trPr()
    before = OxmlElement('w:gridBefore')
    before.set(qn('w:val'), '1')
    props.append(before)
    book = imported(document, tmp_path)
    assert [(b['text'], b['cell']) for b in book['blocks']] == [('Column two', 2), ('Column three', 3)]


@pytest.mark.parametrize('prior', ['ordinary', 'interrupted'])
def test_invalid_vertical_merge_cannot_consume_an_unrelated_cell(tmp_path, prior):
    document = Document()
    table = document.add_table(rows=3, cols=1)
    for i in range(3):
        table.cell(i, 0).text = f'Cell {i}'
    if prior == 'interrupted':
        restart = OxmlElement('w:vMerge')
        restart.set(qn('w:val'), 'restart')
        table._tbl.tr_lst[0].tc_lst[0].get_or_add_tcPr().append(restart)
    merge = OxmlElement('w:vMerge')
    table._tbl.tr_lst[2].tc_lst[0].get_or_add_tcPr().append(merge)
    source = tmp_path / 'invalid-merge.docx'
    document.save(source)
    with pytest.raises(ValueError, match='vertical merge'):
        read_docx(str(source), tmp_path / 'assets')


@pytest.mark.parametrize('count', [2, 12])
def test_distinct_hyperlinks_keep_their_own_targets(tmp_path, count):
    document = Document()
    paragraph = document.add_paragraph()
    for index in range(count):
        node = OxmlElement('w:hyperlink')
        target = f'https://example.org/{index}'
        node.set(qn('r:id'), document.part.relate_to(target, RT.HYPERLINK, is_external=True))
        run = paragraph.add_run(f'Link{index}')
        node.append(run._r)
        paragraph._p.append(node)
        paragraph.add_run(' ')
    book = imported(document, tmp_path)
    assert book['blocks'][0]['links'] == [
        {'text': f'Link{i}', 'href': f'https://example.org/{i}'} for i in range(count)]


def test_failed_asset_replacement_leaves_no_partial_published_asset(tmp_path, monkeypatch):
    import os

    document = Document()
    document.add_paragraph().add_run().add_picture(io.BytesIO(png_bytes(20, 10)))
    source = tmp_path / 'source.docx'
    document.save(source)
    assets = tmp_path / 'assets'
    assets.mkdir()
    (assets / 'keep.bin').write_bytes(b'previous asset')

    def fail_replace(*_args):
        raise OSError('injected replacement failure')
    monkeypatch.setattr(os, 'replace', fail_replace)
    with pytest.raises(OSError, match='injected'):
        read_docx(str(source), assets)
    assert {p.name: p.read_bytes() for p in assets.iterdir()} == {'keep.bin': b'previous asset'}


@pytest.mark.parametrize('seed', range(12))
def test_seeded_mixed_run_events_preserve_the_complete_source_sequence(tmp_path, seed):
    import random
    rng = random.Random(seed)
    document = Document()
    paragraph = document.add_paragraph()
    run = paragraph.add_run()
    note_part(document)
    expected, buffer, count = [], [], 0

    def flush():
        text = ''.join(buffer).strip(' ')
        if text.strip():
            expected.append(('paragraph', text))
        buffer.clear()

    for index in range(18):
        if rng.randrange(3) == 0:
            run = paragraph.add_run()
        event = rng.choice(['text', 'note', 'image', 'tab', 'line', 'page'])
        if event == 'note':
            count += 1
            reference(run)
            buffer.append(f'[[fn:fn{count:04d}]]')
        elif event in {'image', 'page'}:
            flush()
            if event == 'image':
                run.add_picture(io.BytesIO(png_bytes(20, 10)))
                expected.append(('image', ''))
            else:
                run.add_break(WD_BREAK.PAGE)
                expected.append(('pagebreak', ''))
        else:
            text = {'text': f'word{index}', 'tab': '\t', 'line': '\n'}[event]
            run.add_text(text)
            buffer.append(text)
    flush()
    book = imported(document, tmp_path)
    assert [(b['type'], b.get('text') or '') for b in book['blocks']] == expected
    assert len(book['footnotes']) == count
