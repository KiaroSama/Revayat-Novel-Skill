"""Script-applicable emphasis at actual native intake/publication boundaries."""
import argparse
import zipfile
from xml.etree import ElementTree as ET

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE

import bookir as ir
import opc
from build_docx import Builder, add_arguments
from read_docx import read_docx


@pytest.mark.parametrize('script,text,expected', [
    ('complex', 'فارسی', '**فارسی**'), ('complex', 'Latin', 'Latin'),
    ('ordinary', 'Latin', '**Latin**'), ('ordinary', 'فارسی', 'فارسی')])
def test_source_emphasis_uses_applicable_script_property(tmp_path, script, text, expected):
    source = Document()
    run = source.add_paragraph().add_run(text)
    if script == 'complex':
        run.font.cs_bold = True
    else:
        run.bold = True
    path = tmp_path / 'source.docx'
    source.save(path)
    book = read_docx(str(path), tmp_path / 'assets')
    assert book['blocks'][0]['text'] == expected


def test_inherited_complex_emphasis_and_direct_false_override(tmp_path):
    source = Document()
    base = source.styles.add_style('Native Base', WD_STYLE_TYPE.PARAGRAPH)
    base.font.cs_bold = True
    base.font.cs_italic = True
    child = source.styles.add_style('Native Child', WD_STYLE_TYPE.PARAGRAPH)
    child.base_style = base
    paragraph = source.add_paragraph(style=child)
    paragraph.add_run('پررنگ')
    plain = paragraph.add_run(' ساده')
    plain.font.cs_bold = False
    plain.font.cs_italic = False
    path = tmp_path / 'source.docx'
    source.save(path)
    book = read_docx(str(path), tmp_path / 'assets')
    assert book['blocks'][0]['text'] == '***پررنگ*** ساده'


def test_mixed_source_run_does_not_or_complex_emphasis_into_latin(tmp_path):
    source = Document()
    run = source.add_paragraph().add_run('فارسی Latin')
    run.font.cs_bold = True
    run.bold = False
    path = tmp_path / 'mixed.docx'
    source.save(path)
    book = read_docx(str(path), tmp_path / 'assets')
    assert book['blocks'][0]['text'] == '**فارسی** Latin'


def test_mixed_native_note_uses_the_same_script_applicable_emphasis(tmp_path):
    from tests.test_docx_event_order import note_part, reference

    source = Document()
    reference(source.add_paragraph('Text').add_run())
    note_part(source, bodies=[('1', '<w:p><w:r><w:rPr><w:bCs/><w:b w:val="0"/></w:rPr>'
                              '<w:t>فارسی Latin</w:t></w:r></w:p>')])
    path = tmp_path / 'notes.docx'
    source.save(path)
    book = read_docx(str(path), tmp_path / 'assets')
    assert book['footnotes'][0]['text'] == '**فارسی** Latin'


def test_published_body_and_running_runs_write_complex_false_overrides(tmp_path):
    source = Document()
    source.styles['Normal'].font.cs_bold = True
    source.styles['Normal'].font.cs_italic = True
    template = tmp_path / 'template.docx'
    source.save(template)
    book = ir.new_book()
    book['blocks'] = [ir.make_block('paragraph', 1, text='Source', target='**پررنگ** ساده *کج*')]
    parser = argparse.ArgumentParser()
    add_arguments(parser)
    options = parser.parse_args(['--book', 'unused', '--out', 'unused', '--no-toc',
                                 '--template', str(template), '--no-page-numbers'])
    path = tmp_path / 'published.docx'
    Builder(book, tmp_path, options).build(path)
    with zipfile.ZipFile(path) as archive:
        root = ET.fromstring(archive.read('word/document.xml'))
    runs = [r for r in root.iter(opc.qname('w', 'r')) if opc.text_of(r).strip()]
    properties = [r.find(opc.qname('w', 'rPr')) for r in runs]
    # A present bare on/off element is true; python-docx elides its default.
    assert [(p.find(opc.qname('w', 'bCs')).get(opc.qname('w', 'val'), '1'),
             p.find(opc.qname('w', 'iCs')).get(opc.qname('w', 'val'), '1')) for p in properties] == [
        ('1', '0'), ('0', '0'), ('0', '1')]
