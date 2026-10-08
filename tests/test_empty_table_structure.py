"""Declared native table shape survives without invented translation prose."""
import argparse
import copy
import json

import pytest
from docx import Document

import bookir as ir
import published
from build_docx import Builder, add_arguments
from read_docx import read_docx
import qa


def roundtrip(source, tmp_path):
    path = tmp_path / 'source.docx'
    source.save(path)
    book = read_docx(str(path), tmp_path / 'assets')
    book['meta'].update(title='', author='')
    for block in ir.iter_text_blocks(book):
        block['target'] = block['text']
    assert ir.validate_book(book) == []
    book = ir.decode_book(json.dumps(book, ensure_ascii=False))
    parser = argparse.ArgumentParser()
    add_arguments(parser)
    options = parser.parse_args(['--book', 'unused', '--out', 'unused', '--no-toc', '--no-page-numbers'])
    output = tmp_path / 'published.docx'
    Builder(book, tmp_path / 'assets', options).build(output)
    return book, Document(output), output


@pytest.mark.parametrize('populated', [False, True])
def test_empty_trailing_table_grid_is_source_authoritative(tmp_path, populated):
    source = Document()
    source.add_paragraph('Before')
    table = source.add_table(rows=2, cols=2)
    if populated:
        table.cell(0, 0).text = 'Only prose'
    source.add_paragraph('After')
    book, output, path = roundtrip(source, tmp_path)
    assert len(output.tables) == 1
    assert (len(output.tables[0].rows), len(output.tables[0].columns)) == (2, 2)
    assert len(list(ir.iter_text_blocks(book))) == (3 if populated else 2)
    assert len([b for b in book['blocks'] if b['type'] == 'table']) == 1
    assert qa.check_docx(path, book).summary()['ok']
    changed = copy.deepcopy(book)
    changed['tables'][0]['columns'] = 3
    assert published.structure(changed) != published.structure(book)
    assert not qa.check_docx(path, changed).summary()['ok']


def test_child_only_table_keeps_immediate_empty_parent_cell(tmp_path):
    source = Document()
    outer = source.add_table(rows=2, cols=2)
    outer.cell(0, 1).add_table(rows=1, cols=1).cell(0, 0).text = 'Inside'
    book, output, path = roundtrip(source, tmp_path)
    assert len(output.tables) == 1
    outer = output.tables[0]
    assert len(outer.rows) == len(outer.columns) == 2
    assert len(outer.cell(0, 1).tables) == 1
    assert outer.cell(0, 1).tables[0].cell(0, 0).text == 'Inside'
    assert qa.check_docx(path, book).summary()['ok']


def test_empty_merged_edges_keep_source_spans(tmp_path):
    source = Document()
    table = source.add_table(rows=2, cols=3)
    table.cell(0, 1).merge(table.cell(1, 2))
    book, output, path = roundtrip(source, tmp_path)
    rebuilt = output.tables[0]
    assert len(rebuilt.rows) == 2 and len(rebuilt.columns) == 3
    assert rebuilt.cell(0, 1)._tc is rebuilt.cell(1, 2)._tc
    assert any(c['row_span'] == 2 and c['col_span'] == 2 for c in book['tables'][0]['cells'])
    assert qa.check_docx(path, book).summary()['ok']


def test_authored_empty_cell_paragraphs_survive_apart_from_terminator(tmp_path):
    source = Document()
    cell = source.add_table(rows=1, cols=1).cell(0, 0)
    cell.paragraphs[0].text = 'Before'
    cell.add_paragraph('')
    cell.add_paragraph('After')
    cell.add_paragraph('')
    cell.add_paragraph('')
    book, output, path = roundtrip(source, tmp_path)
    assert [b['controls'] for b in book['blocks'] if b['type'] == 'layout'] == ['', '']
    assert [p.text for p in output.tables[0].cell(0, 0).paragraphs] == ['Before', '', 'After', '', '']
    assert qa.check_docx(path, book).summary()['ok']
    removed = Document(path)
    blank = removed.tables[0].cell(0, 0).paragraphs[1]._p
    blank.getparent().remove(blank)
    removed.save(path)
    assert not qa.check_docx(path, book).summary()['ok']


def test_empty_table_moved_across_body_text_fails_native_package_order(tmp_path):
    import zipfile
    from xml.etree import ElementTree as ET
    import opc

    source = Document()
    source.add_paragraph('Before')
    source.add_table(rows=1, cols=1)
    source.add_paragraph('After')
    book, _, path = roundtrip(source, tmp_path)
    assert qa.check_docx(path, book).summary()['ok']
    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    root = ET.fromstring(parts['word/document.xml'])
    body = root.find(opc.qname('w', 'body'))
    table = body.find(opc.qname('w', 'tbl'))
    body.remove(table)
    body.insert(len(body) - 1, table)
    parts['word/document.xml'] = ET.tostring(root, encoding='utf-8')
    with zipfile.ZipFile(path, 'w') as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    assert 'body-structure-order-mismatch' in {f['code'] for f in qa.check_docx(path, book).findings}


@pytest.mark.parametrize('moving', ['layout', 'table'])
def test_structural_events_cannot_move_across_an_image_anchor(tmp_path, moving):
    import io
    import zipfile
    from xml.etree import ElementTree as ET
    import opc
    from tests_support import png_bytes

    source = Document()
    source.add_paragraph('Before')
    source.add_paragraph('')
    source.add_table(rows=1, cols=1)
    source.add_paragraph().add_run().add_picture(io.BytesIO(png_bytes(10, 10)))
    source.add_paragraph('After')
    book, _, path = roundtrip(source, tmp_path)
    assert qa.check_docx(path, book).summary()['ok']
    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    root = ET.fromstring(parts['word/document.xml'])
    body = root.find(opc.qname('w', 'body'))
    event = body.find(opc.qname('w', 'tbl')) if moving == 'table' else next(
        node for node in body if node.tag == opc.qname('w', 'p') and not opc.text_of(node)
        and not list(node.iter(opc.qname('w', 'drawing'))))
    body.remove(event)
    body.insert(len(body) - 1, event)
    parts['word/document.xml'] = ET.tostring(root, encoding='utf-8')
    with zipfile.ZipFile(path, 'w') as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    assert 'body-structure-order-mismatch' in {f['code'] for f in qa.check_docx(path, book).findings}


def test_native_note_and_metadata_with_table_are_valid_order_anchors(tmp_path):
    from tests.test_docx_event_order import note_part, reference

    source = Document()
    reference(source.add_paragraph('Before ').add_run())
    note_part(source)
    source.add_table(rows=1, cols=1)
    source.add_paragraph('After')
    book, _, _ = roundtrip(source, tmp_path)
    book['meta'].update(title='Title', author='Author')
    parser = argparse.ArgumentParser()
    add_arguments(parser)
    options = parser.parse_args(['--book', 'unused', '--out', 'unused', '--no-toc', '--no-page-numbers'])
    path = tmp_path / 'metadata.docx'
    Builder(book, tmp_path / 'assets', options).build(path)
    assert qa.check_docx(path, book).summary()['ok']


@pytest.mark.parametrize('defect', ['bounds', 'overlap', 'cycle', 'event-order', 'terminal-type'])
def test_invalid_structure_refuses_persisted_book(tmp_path, defect):
    source = Document()
    source.add_table(rows=2, cols=2)
    book, _, _ = roundtrip(source, tmp_path)
    record = book['tables'][0]
    if defect == 'bounds':
        record['cells'][0]['row_span'] = 3
    elif defect == 'overlap':
        record['cells'].append(dict(record['cells'][0]))
    elif defect == 'cycle':
        record.update(parent_table=record['id'], parent_row=1, parent_cell=1)
    elif defect == 'terminal-type':
        record['cells'][0]['terminal_empty'] = 'true'
    else:
        book['blocks'] = [b for b in book['blocks'] if b['type'] != 'table']
    with pytest.raises(ValueError, match='structure|table'):
        ir.decode_book(json.dumps(book, ensure_ascii=False))
