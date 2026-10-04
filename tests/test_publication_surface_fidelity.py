"""Published Word surfaces preserve authored controls and source anchor semantics."""

from __future__ import annotations

import argparse
import copy
import logging
import zipfile
from xml.etree import ElementTree as ET

import pytest

import bookir as ir
import opc
import qa
from build_docx import Builder, add_arguments
from tests_support import png_bytes

LOG = logging.getLogger(__name__)


def publish(book, work, *, toc=False):
    parser = argparse.ArgumentParser()
    add_arguments(parser)
    opts = parser.parse_args(
        [
            "--book",
            "unused",
            "--out",
            "unused",
            "--font",
            "DejaVu Sans",
            "--no-justify",
            "--page-breaks",
            "source",
        ]
    )
    opts.toc = toc
    path = work / "published.docx"
    report = Builder(copy.deepcopy(book), work, opts).build(path)
    with zipfile.ZipFile(path) as archive:
        parts = {
            name: ET.fromstring(archive.read(name))
            for name in archive.namelist()
            if name.startswith("word/") and name.endswith(".xml")
        }
    LOG.info(
        "Published synthetic fidelity fixture with %d warnings", report["warning_count"]
    )
    return path, parts, report


def fixture():
    book = ir.new_book(source_format="docx", pages=1)
    book["meta"].update(title="", author="")
    book["blocks"] = [
        ir.make_block("paragraph", 1, page=1, text="Source", target="متن")
    ]
    return book


@pytest.mark.parametrize("edge", ["\n", "\n\n", "\t", " \t", "\t\n"])
@pytest.mark.parametrize(
    "surface", ["body", "cell", "note", "header", "caption", "title", "byline"]
)
def test_authored_edge_controls_reach_their_actual_word_surface(
    tmp_path, edge, surface
):
    book = fixture()
    target = edge + "متن مشخص" + edge
    block = book["blocks"][0]
    if surface in {"body", "cell"}:
        block["target"] = target
        if surface == "cell":
            block.update(table="tbl0001", row=1, cell=1)
    elif surface == "note":
        block.update(text="Source[[fn:fn0001]]", target="متن[[fn:fn0001]]")
        note = ir.make_footnote(1, anchor_block=block["id"], text="Source note")
        note["target"] = target
        book["footnotes"] = [note]
    elif surface == "header":
        book["sections"] = [
            {
                "start_block": block["id"],
                "headers": {
                    "default": {
                        "paragraphs": [
                            {
                                "pieces": [
                                    {
                                        "id": "rh0001",
                                        "text": "Source header",
                                        "target": target,
                                    }
                                ]
                            }
                        ]
                    }
                },
            }
        ]
    elif surface == "caption":
        data = png_bytes(20, 10)
        (tmp_path / "image.png").write_bytes(data)
        book["blocks"].append(
            ir.make_block(
                "image",
                2,
                page=1,
                asset="image.png",
                sha256=ir.sha256_bytes(data),
                width_pt=20,
                height_pt=10,
                alt="Source caption",
                target_alt=target,
            )
        )
    elif surface == "title":
        book["meta"].update(title="Source title", title_target=target)
    else:
        book["meta"].update(
            title="Source title",
            title_target="عنوان",
            author="Source author",
            author_target=target,
        )
    path, parts, _ = publish(book, tmp_path)
    if surface == "note":
        assert list(opc.footnote_bodies(parts["word/footnotes.xml"]).values()) == [
            target
        ]
    elif surface == "header":
        assert target in [
            opc.text_of(p)
            for name, root in parts.items()
            if name.startswith("word/header")
            for p in root.iter(opc.qname("w", "p"))
        ]
    else:
        assert target in [
            opc.text_of(p) for p in parts["word/document.xml"].iter(opc.qname("w", "p"))
        ]
    assert qa.check_docx(path, book).summary()["ok"]


def test_author_only_metadata_is_published_without_a_title(tmp_path):
    book = fixture()
    book["meta"].update(author="Writer", author_target="نویسنده")
    _, parts, _ = publish(book, tmp_path)
    assert "نویسنده" in opc.text_of(parts["word/document.xml"])


@pytest.mark.parametrize("nested_depth", [1, 2, 3])
def test_every_cell_ends_in_a_paragraph_after_a_terminal_nested_table(
    tmp_path, nested_depth
):
    book = fixture()
    book["blocks"] = []
    for depth in range(nested_depth + 1):
        fields = dict(table=f"tbl{depth + 1:04d}", row=1, cell=1)
        if depth:
            fields.update(parent_table=f"tbl{depth:04d}", parent_row=1, parent_cell=1)
        book["blocks"].append(
            ir.make_block(
                "paragraph",
                depth + 1,
                page=1,
                text=f"Layer {depth}",
                target=f"Layer {depth}",
                **fields,
            )
        )
    path, parts, _ = publish(book, tmp_path)
    cells = list(parts["word/document.xml"].iter(opc.qname("w", "tc")))
    assert len(cells) == nested_depth + 1
    assert all(list(cell)[-1].tag == opc.qname("w", "p") for cell in cells)
    assert qa.check_docx(path, book).summary()["ok"]


@pytest.mark.parametrize(
    "kind, expected_style",
    [
        ("heading", "Heading2"),
        ("blockquote", "Quote"),
        ("caption", "Caption"),
        ("listitem", "ListBullet"),
    ],
)
def test_semantic_cell_blocks_keep_the_same_style_as_body_blocks(
    tmp_path, kind, expected_style
):
    book = fixture()
    book["blocks"][0].update(type=kind, level=2, table="tbl0001", row=1, cell=1)
    path, parts, _ = publish(book, tmp_path, toc=True)
    paragraph = next(parts["word/document.xml"].iter(opc.qname("w", "tc"))).find(
        opc.qname("w", "p")
    )
    assert (
        paragraph.find(opc.qname("w", "pPr"))
        .find(opc.qname("w", "pStyle"))
        .get(opc.qname("w", "val"))
        == expected_style
    )
    assert qa.check_docx(path, book).summary()["ok"]


@pytest.mark.parametrize("kind", ["image", "separator", "pagebreak"])
@pytest.mark.parametrize("in_cell", [False, True])
def test_source_anchors_survive_nonprose_body_and_cell_events(tmp_path, kind, in_cell):
    book = fixture()
    marker = ir.make_block(kind, 2, page=1, soft=False, bookmarks=["destination"])
    if in_cell:
        marker.update(table="tbl0001", row=1, cell=1)
    if kind == "image":
        data = png_bytes(20, 10)
        (tmp_path / "image.png").write_bytes(data)
        marker.update(
            asset="image.png", sha256=ir.sha256_bytes(data), width_pt=20, height_pt=10
        )
    book["blocks"] += [
        marker,
        ir.make_block(
            "paragraph",
            3,
            page=1,
            text="Go",
            target="Go",
            links=[{"text": "Go", "href": "#destination"}],
        ),
    ]
    path, parts, _ = publish(book, tmp_path)
    assert any(
        n.get(opc.qname("w", "name")) == "destination"
        for n in parts["word/document.xml"].iter(opc.qname("w", "bookmarkStart"))
    )
    if kind == "separator":
        assert "❖" in opc.text_of(parts["word/document.xml"])
    assert qa.check_docx(path, book).summary()["ok"]


def test_package_gate_rejects_a_removed_cell_terminal_paragraph(tmp_path):
    book = fixture()
    book["blocks"][0].update(table="tbl0001", row=1, cell=1)
    path, _, _ = publish(book, tmp_path)
    with zipfile.ZipFile(path) as z:
        contents = {name: z.read(name) for name in z.namelist()}
    root = ET.fromstring(contents["word/document.xml"])
    cell = next(root.iter(opc.qname("w", "tc")))
    cell.remove(list(cell)[-1])
    contents["word/document.xml"] = ET.tostring(root, encoding="utf-8")
    with zipfile.ZipFile(path, "w") as z:
        for name, data in contents.items():
            z.writestr(name, data)
    report = qa.check_docx(path, book).summary()
    assert any(
        f["code"] == "table-cell-terminal-paragraph" for f in report["findings"]
    ), report


@pytest.mark.parametrize("label", ["\nGo\n", "\tGo\t", " Go ", "\t\nGo\n\t"])
def test_link_labels_keep_authored_edge_controls(tmp_path, label):
    book = fixture()
    book["blocks"][0].update(
        text=label,
        target=label,
        links=[{"text": label, "href": "https://example.org/path"}],
    )
    path, parts, _ = publish(book, tmp_path)
    links = list(parts["word/document.xml"].iter(opc.qname("w", "hyperlink")))
    assert len(links) == 1
    assert opc.text_of(links[0]) == label
    assert qa.check_docx(path, book).summary()["ok"]


@pytest.mark.parametrize("corruption", [False, True])
def test_bookmark_checks_are_namespace_based_not_prefix_based(tmp_path, corruption):
    book = fixture()
    book["blocks"][0].update(type="heading", level=2, bookmarks=["destination"])
    book["blocks"].append(
        ir.make_block(
            "paragraph",
            2,
            text="Go",
            target="Go",
            links=[{"text": "Go", "href": "#destination"}],
        )
    )
    path, _, _ = publish(book, tmp_path, toc=True)
    with zipfile.ZipFile(path) as z:
        contents = {name: z.read(name) for name in z.namelist()}
    # ElementTree serializes legal namespace aliases rather than the original w prefix.
    root = ET.fromstring(contents["word/document.xml"])
    if corruption:
        for parent in root.iter():
            for node in list(parent):
                if (
                    node.tag == opc.qname("w", "bookmarkStart")
                    and node.get(opc.qname("w", "name")) == "destination"
                ):
                    parent.remove(node)
    contents["word/document.xml"] = ET.tostring(root, encoding="utf-8")
    with zipfile.ZipFile(path, "w") as z:
        for name, data in contents.items():
            z.writestr(name, data)
    report = qa.check_docx(path, book).summary()
    if corruption:
        assert any(f["code"] == "dead-link" for f in report["findings"]), report
    else:
        assert report["ok"], report


def native_pipeline(work, route):
    """Produce a real reader/worksheet/merge/build fixture for tests and rendering."""
    import chunk
    import merge
    import pagerun
    import read_docx as reader
    import worksheet
    from docx import Document
    from test_docx_event_order import note_part, reference

    work.mkdir(parents=True, exist_ok=True)
    source = Document()
    source.core_properties.title = ""
    source.core_properties.author = ""
    cells = source.add_table(rows=1, cols=2).rows[0].cells
    cells[0].paragraphs[0].style = "Heading 2"
    cells[0].paragraphs[0].text = "عنوان جدول"
    body = cells[0].add_paragraph("\nآغاز بند\n\nپایان بند\n")
    reference(body.add_run())
    cells[0].add_table(rows=1, cols=1).cell(0, 0).text = "جدول درونی"
    cells[1].paragraphs[0].text = "ستون دوم"
    note_part(
        source,
        bodies=[
            (
                "1",
                "<w:p><w:r><w:br/><w:t>خط اول</w:t><w:tab/>"
                "<w:t>یادداشت</w:t><w:br/><w:t>خط دوم</w:t><w:br/></w:r></w:p>",
            )
        ],
    )
    src = work / "source.docx"
    source.save(src)
    book = reader.read_docx(str(src), work / "assets")
    book["meta"].update(title="", author="")
    for b in book["blocks"]:
        b["page"] = 1
    book["pages"] = 1
    path, jobs = work / "book.json", work / route
    ir.save_book(book, path)
    manifest = (chunk if route == "chunks" else pagerun).build(
        path, jobs, glossary_path=None, budget=5000
    )
    for entry in manifest["chunks"]:
        lines = [worksheet.request_line(entry["request"])]
        for unit in worksheet.read_worksheet(
            (jobs / entry["file"]).read_text(encoding="utf-8")
        ):
            lines += [
                f"@@ {unit['id']} {unit['kind']}",
                worksheet.escape_payload(unit["text"]),
            ]
        ir.write_text(jobs / entry["output"], "\n".join(lines) + "\n")
    result = (
        merge.merge(path, jobs)
        if route == "chunks"
        else pagerun.merge_page(path, jobs, 1)
    )
    assert result["ok"], result
    translated = ir.load_book(path)
    destination, parts, report = publish(translated, work, toc=True)
    assert qa.check_docx(destination, translated).summary()["ok"]
    bodyxml = parts["word/document.xml"]
    for cell in bodyxml.iter(opc.qname("w", "tc")):
        assert cell[-1].tag == opc.qname("w", "p")
    heading = next(bodyxml.iter(opc.qname("w", "tbl"))).find(
        ".//" + opc.qname("w", "p")
    )
    assert (
        heading.find(".//" + opc.qname("w", "pStyle")).get(opc.qname("w", "val"))
        == "Heading2"
    )
    assert any(
        n.get(opc.qname("w", "name"), "").startswith("rv_") for n in heading.iter()
    )
    assert list(opc.footnote_bodies(parts["word/footnotes.xml"]).values()) == [
        "\nخط اول\tیادداشت\nخط دوم\n"
    ]
    assert any(
        opc.text_of(p).startswith("\nآغاز") for p in bodyxml.iter(opc.qname("w", "p"))
    )
    return destination


@pytest.mark.parametrize("route", ["chunks", "pages"])
def test_native_table_and_note_controls_survive_both_production_routes(tmp_path, route):
    native_pipeline(tmp_path, route)
