"""Native source order survives both worksheet routes and the real Word builder."""
from __future__ import annotations

import argparse
import copy
import io
import logging
import zipfile

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
import pytest

import bookir as ir
import chunk
import merge
import opc
import pagerun
import qa
import worksheet
from build_docx import Builder, add_arguments
from read_docx import read_docx
from test_docx_event_order import note_part, reference
from tests_support import png_bytes

LOG = logging.getLogger(__name__)


@pytest.mark.parametrize('route', ['chunks', 'pages'])
def test_native_note_and_drawing_sequence_reaches_the_final_package(tmp_path, route):
    document = Document()
    document.core_properties.title = 'Intake fixture'
    document.core_properties.author = ''
    document.add_heading('Opening', level=1)
    run = document.add_paragraph().add_run('Before ')
    reference(run)
    run.add_text(' after note.\n\nBefore image. ')
    image = png_bytes(120, 80)
    run.add_picture(io.BytesIO(image), width=Pt(90))
    run.add_text('After image.')
    note_part(document, bodies=[('1',
        '<w:p><w:r><w:t>First</w:t><w:tab/><w:t>column</w:t><w:br/><w:t>Second line.</w:t></w:r></w:p>')])
    source = tmp_path / 'source.docx'
    document.save(source)
    assets = tmp_path / 'assets'
    book = read_docx(str(source), assets)
    before = copy.deepcopy(book)
    path, jobs = tmp_path / 'book.json', tmp_path / route
    ir.save_book(book, path)
    manifest = (chunk if route == 'chunks' else pagerun).build(path, jobs, glossary_path=None)
    for entry in manifest['chunks']:
        units = worksheet.read_worksheet((jobs / entry['file']).read_text(encoding='utf-8'))
        reply = [worksheet.request_line(entry['request'])]
        for unit in units:
            # This is transport evidence, not a claim of literary translation.
            reply += [f"@@ {unit['id']} {unit['kind']}", worksheet.escape_payload('متن آزمایشی. ' + unit['text'])]
        ir.write_text(jobs / entry['output'], '\n'.join(reply) + '\n')
    if route == 'chunks':
        assert merge.merge(path, jobs)['ok']
    else:
        for page in sorted({entry['page'] for entry in manifest['chunks']}):
            result = pagerun.merge_page(path, jobs, page)
            assert result['ok'], result
    after = ir.load_book(path)
    assert [b.get('text') for b in before['blocks']] == [b.get('text') for b in after['blocks']]
    assert after['footnotes'][0]['text'] == 'First\tcolumn\nSecond line.'
    parser = argparse.ArgumentParser()
    add_arguments(parser)
    destination = tmp_path / 'out.docx'
    opts = parser.parse_args(['--book', str(path), '--out', str(destination), '--font', 'DejaVu Sans'])
    Builder(after, assets, opts).build(destination)
    report = qa.check_docx(destination, after).summary()
    assert report['ok'], report
    with zipfile.ZipFile(destination) as archive:
        from lxml import etree
        xml = etree.fromstring(archive.read('word/document.xml'))
        events = []
        for node in xml.iter():
            if node.tag == qn('w:t'):
                events.append(node.text or '')
            elif node.tag == qn('w:footnoteReference'):
                events.append('[NOTE]')
            elif node.tag == qn('w:drawing'):
                events.append('[IMAGE]')
        text = ''.join(events)
        assert text.index('Before ') < text.index('[NOTE]') < text.index(' after note.')
        assert text.index('Before image.') < text.index('[IMAGE]') < text.index('After image.')
        notes = etree.fromstring(archive.read('word/footnotes.xml'))
        assert len(list(notes.iter(qn('w:tab')))) == 1
        assert len(list(notes.iter(qn('w:br')))) == 1
        assert opc.text_of(notes).find('Second line.') > opc.text_of(notes).find('First')
    image_block = next(b for b in after['blocks'] if b['type'] == 'image')
    assert (assets / image_block['asset']).read_bytes() == image
    LOG.info('Verified %s native DOCX pipeline and package event order', route)
