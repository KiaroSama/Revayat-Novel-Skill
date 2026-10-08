"""Authored layout is ordered native content with zero prose units."""
import argparse

from docx import Document

import bookir as ir
import published
from build_docx import Builder, add_arguments
from read_docx import read_docx
from tests.test_epub_source_integrity import make_epub


def test_native_control_only_paragraphs_publish_without_translation_units(tmp_path):
    source = Document()
    for text in ['Before', '\t\n\n', '', 'After']:
        source.add_paragraph(text)
    path = tmp_path / 'source.docx'
    source.save(path)
    book = read_docx(str(path), tmp_path / 'assets')
    book['meta'].update(title='', author='')
    assert [b['controls'] for b in book['blocks'] if b['type'] == 'layout'] == ['\t\n\n', '']
    assert [u['id'] for u in published.units(book)] == [b['id'] for b in ir.iter_text_blocks(book)]
    for block in ir.iter_text_blocks(book):
        block['target'] = block['text']
    parser = argparse.ArgumentParser()
    add_arguments(parser)
    options = parser.parse_args(['--book', 'unused', '--out', 'unused', '--no-toc', '--no-page-numbers'])
    out = tmp_path / 'out.docx'
    Builder(book, tmp_path, options).build(out)
    assert [p.text for p in Document(out).paragraphs] == ['Before', '\t\n\n', '', 'After']


def test_saved_web_break_only_paragraph_reaches_native_epub_layout(tmp_path):
    from webhtml import chapter_html
    from read_epub import read_epub

    raw = '<article><p>Before</p><p><br/><br/></p><p>After</p></article>'
    normalized = chapter_html(raw, 'article', 'Chapter', lambda source: source)
    path = make_epub(tmp_path, {'one.xhtml': normalized.decode('utf-8')})
    book = read_epub(str(path), tmp_path / 'assets')
    assert [b['controls'] for b in book['blocks'] if b['type'] == 'layout'] == ['\n\n']


def test_epub_controls_do_not_claim_pending_empty_anchor(tmp_path):
    from read_epub import read_epub

    path = make_epub(tmp_path, {'one.xhtml': '<p>Before <a href="#dest">jump</a></p>'
                          '<span id="dest"></span><p><br/><br/></p><p>After</p>'})
    book = read_epub(str(path), tmp_path / 'assets')
    layouts = [b for b in book['blocks'] if b['type'] == 'layout']
    assert [b['controls'] for b in layouts] == ['\n\n']
    assert not any(b.get('bookmarks') for b in layouts)
    assert next(b for b in book['blocks'] if b.get('text') == 'After')['bookmarks'] == ['dest']
