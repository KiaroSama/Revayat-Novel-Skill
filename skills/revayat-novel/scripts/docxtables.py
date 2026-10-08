"""Native table topology, independent of whether cells contain prose."""
from __future__ import annotations
from typing import Any
from docx.oxml.ns import qn

def column_span(tc) -> int:
    """How many grid columns this ``w:tc`` covers, from ``w:gridSpan``."""
    properties = tc.find(qn("w:tcPr"))
    if properties is None:
        return 1
    grid = properties.find(qn("w:gridSpan"))
    if grid is None:
        return 1
    try:
        width = int(grid.get(qn("w:val")))
    except (TypeError, ValueError) as error:
        raise ValueError("invalid DOCX table grid span") from error
    if width < 1:
        raise ValueError("invalid DOCX table grid span")
    return width


def _vertical_merge(tc) -> str | None:
    """``"restart"``, ``"continue"``, or ``None`` when the cell is not merged."""
    properties = tc.find(qn("w:tcPr"))
    if properties is None:
        return None
    merge = properties.find(qn("w:vMerge"))
    if merge is None:
        return None
    value = merge.get(qn('w:val'), 'continue')
    if value not in {'restart', 'continue'}:
        raise ValueError('invalid DOCX table vertical merge value')
    return value


def _continuation_is_empty(tc) -> bool:
    """Allow generated empty paragraphs, never discard authored live content."""
    for child in tc:
        if child.tag == qn("w:tcPr"):
            continue
        if child.tag != qn("w:p"):
            return False
        for item in child:
            if item.tag == qn("w:pPr"):
                continue
            if item.tag != qn("w:r"):
                return False
            for part in item:
                if part.tag == qn("w:rPr"):
                    continue
                if part.tag != qn("w:t") or part.text:
                    return False
    return True


def table_cells(table) -> list[dict[str, Any]]:
    """Every distinct cell of a table, once, with its position and span.

    Walks ``w:tr``/``w:tc`` rather than ``row.cells``. Two reasons, and the
    second one is the reason this is not shorter:

    ``row.cells`` *expands* merges — a cell spanning two columns comes back
    twice, and a vertically merged one comes back on every row it covers. Since
    every cell here becomes its own worksheet unit, that made a merged cell's
    sentence get translated twice and printed twice.

    And de-duplicating that expansion by object identity does not work.
    `cell._tc` hands back a fresh lxml proxy on each access, and CPython reuses
    the `id()` of a freed one — measured on a 3x3 grid, three different cells
    reported the same id and a fourth reported one that had never been seen.
    Walking the XML gives each cell exactly once by construction, so there is
    nothing to de-duplicate.

    Returns ``{"tc", "row", "cell", "row_span", "col_span"}`` with 1-based
    ``row``/``cell`` counted in grid columns.
    """
    from docx.table import _Cell

    found: list[dict[str, Any]] = []
    # Grid column -> the record of the cell currently open there, so a
    # `continue` row extends the span of the cell that started it.
    open_at: dict[int, dict[str, Any]] = {}

    for row_number, tr in enumerate(table._tbl.findall(qn("w:tr")), start=1):
        properties = tr.find(qn("w:trPr"))
        before = properties.find(qn("w:gridBefore")) if properties is not None else None
        try:
            column = int(before.get(qn("w:val"))) if before is not None else 0
        except (TypeError, ValueError) as error:
            raise ValueError("invalid DOCX table leading grid offset") from error
        if column < 0:
            raise ValueError("invalid DOCX table leading grid offset")
        following: dict[int, dict[str, Any]] = {}
        for tc in tr.findall(qn("w:tc")):
            width = column_span(tc)
            merge = _vertical_merge(tc)
            if merge == "continue":
                record = open_at.get(column)
                if record is None or record["col_span"] != width:
                    raise ValueError("invalid DOCX table vertical merge continuation")
                if not _continuation_is_empty(tc):
                    raise ValueError("populated DOCX table vertical merge continuation")
                record["row_span"] += 1
                following[column] = record
            else:
                record = {"tc": _Cell(tc, table), "row": row_number,
                          "cell": column + 1, "row_span": 1, "col_span": width}
                found.append(record)
                if merge == "restart":
                    following[column] = record
            column += width
        # Only a restarted/continued span in the immediately preceding row is
        # eligible. Ordinary or absent cells must never seed a later merge.
        open_at = following
    return found


def write_table(writer, blocks, identity, *, container=None):
    """Write declared source shape and native cell events, including empty cells."""
    from docx.oxml import OxmlElement
    import ooxml

    record = next(r for r in writer.book['tables'] if r['id'] == identity)
    rows, columns = record['rows'], record['columns']
    table = (writer.document.add_table(rows=rows, cols=columns) if container is None
             else container.add_table(rows=rows, cols=columns))
    try:
        table.style = writer.document.styles['Table Grid']
    except KeyError:
        pass
    if writer.options.rtl:
        ooxml.set_table_rtl(table)
    cells = {}
    for spec in record['cells']:
        row, column = spec['row'] - 1, spec['cell'] - 1
        cell = table.cell(row, column)
        if spec['row_span'] > 1 or spec['col_span'] > 1:
            cell = cell.merge(table.cell(row + spec['row_span'] - 1,
                                         column + spec['col_span'] - 1))
        cells[(spec['row'], spec['cell'])] = cell
    for cell in cells.values():
        for paragraph in list(cell.paragraphs):
            cell._tc.remove(paragraph._p)
    nested = set()
    for block in blocks:
        if block.get('table') == identity:
            if block['type'] != 'table':
                writer._cell_event(cells[(block['row'], block['cell'])], block)
        elif block.get('parent_table') == identity and block.get('table') not in nested:
            child_id = block['table']
            nested.add(child_id)
            owner = cells[(block['parent_row'], block['parent_cell'])]
            descendants = {child_id}
            for item in blocks:
                if item.get('parent_table') in descendants:
                    descendants.add(item['table'])
            write_table(writer, [b for b in blocks if b.get('table') in descendants], child_id, container=owner)
            # add_table adds a required trailing paragraph. Restore it after all
            # authored following events instead of inserting an extra blank line.
            last = owner._tc[-1]
            if last.tag == qn('w:p') and len(last) == 0:
                owner._tc.remove(last)
    last_events = {(b.get('row'), b.get('cell')): b for b in blocks
                   if b.get('table') == identity and b['type'] != 'table'}
    for spec in record['cells']:
        cell = cells[(spec['row'], spec['cell'])]
        last = last_events.get((spec['row'], spec['cell']), {})
        if spec.get('terminal_empty') or cell._tc[-1].tag != qn('w:p') or (
                last.get('type') == 'layout' and last.get('controls') == ''):
            writer.paragraph(container=cell)
    # Rows may omit leading/trailing grid positions. A source-authoritative
    # inventory is not permission to create empty cells where none existed.
    for number, tr in enumerate(table._tbl.tr_lst, 1):
        covered = sorted({c for spec in record['cells']
                          if spec['row'] <= number < spec['row'] + spec['row_span']
                          for c in range(spec['cell'], spec['cell'] + spec['col_span'])})
        leading = covered[0] - 1 if covered else columns
        trailing = columns - covered[-1] if covered else 0
        cursor = 1
        for tc in list(tr.tc_lst):
            width = column_span(tc)
            if cursor <= leading or cursor > columns - trailing:
                tr.remove(tc)
            cursor += width
        for tag, count in [('gridBefore', leading), ('gridAfter', trailing)]:
            if count:
                node = OxmlElement('w:' + tag)
                node.set(qn('w:val'), str(count))
                tr.get_or_add_trPr().append(node)
    return table
