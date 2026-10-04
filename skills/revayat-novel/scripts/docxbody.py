"""Native publication of ordered body blocks, tables, images and source breaks."""
from __future__ import annotations

from pathlib import Path
from typing import Any

from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Pt

import bookir as ir
import ooxml

ASSUMED_DPI = 96.0
_STYLE_BY_TYPE = {"paragraph": "Normal", "blockquote": "Quote",
                  "caption": "Caption", "verse": "Normal"}

def _has_style(document, name: str) -> bool:
    try:
        document.styles[name]
        return True
    except KeyError:
        return False


class BodyWriter:
    def body(self) -> None:
        first_heading = True
        blocks = self.book.get("blocks", [])
        index = 0
        while index < len(blocks):
            block = blocks[index]
            kind = block["type"]

            # Before anything is written for this block: a section break sits
            # in front of the block that opens the section, not after it.
            record = self.section_at.get(block["id"])
            if record is not None:
                self.start_section(record)

            # A table arrives as a run of consecutive blocks carrying the same
            # `table` id, each tagged with its row and cell. Keeping them as
            # ordinary blocks is what lets every cell be translated as its own
            # worksheet unit; putting the grid back is this builder's job.
            table_id = block.get("table")
            if table_id:
                run = []
                descendants = {table_id}
                while index < len(blocks):
                    item = blocks[index]
                    if item.get("table") not in descendants:
                        if item.get("parent_table") not in descendants:
                            break
                        descendants.add(item["table"])
                    run.append(item)
                    index += 1
                self._table(run)
                continue
            index += 1

            if kind == "heading":
                self._heading(block, first_heading)
                first_heading = False
            elif kind == "image":
                self._image(block)
            elif kind == "pagebreak":
                self._pagebreak(block)
            elif kind == "separator":
                paragraph = self.paragraph("Normal", align=WD_ALIGN_PARAGRAPH.CENTER)
                self.write_markup(paragraph, "❖")
            elif kind in ir.TEXT_TYPES:
                self._text_block(block)

    def _heading(self, block: dict[str, Any], first: bool) -> None:
        level = min(6, max(1, int(block.get("level", 1))))
        paragraph = self.paragraph(
            f"Heading {level}",
            align=WD_ALIGN_PARAGRAPH.CENTER if level == 1 else None,
        )
        if level == 1 and self.options.page_breaks != "none" and not first:
            ooxml.page_break_before(paragraph)
        ooxml.keep_with_next(paragraph)
        anchor = block.get("_anchor")
        text = self._text_of(block)
        self.write_markup(paragraph, text, block.get("links"))
        self._apply_source_size(paragraph, block)
        if anchor:
            self.bookmarks.wrap(paragraph, anchor, ir.plain_text(text))
        self._place_bookmarks(paragraph, block)

    def _apply_source_size(self, paragraph, block: dict[str, Any]) -> None:
        """Reproduce the heading's size from the source book, when asked.

        Extraction records the real point size of every heading it finds. The
        default still uses the Word ``Heading N`` styles — they give a coherent
        document even when the source's own sizes are erratic — but
        ``--heading-size source`` reinstates the book's exact metrics, which is
        what "keep the book's styling" actually means.
        """
        if self.options.heading_size != "source":
            return
        size = block.get("font_size_pt")
        if not size:
            return
        for run in paragraph.runs:
            run.font.size = Pt(float(size))

    def _table(self, cells: list[dict[str, Any]], *, container=None) -> None:
        """Rebuild one table from the blocks its cells were flattened into.

        The cells were kept as ordinary text blocks all the way through the
        pipeline on purpose: one block is one worksheet unit, so every cell gets
        translated, counted and gated exactly like a paragraph. Only the grid is
        missing by then, and only here is there anything to put it back into.

        A cell that lost its coordinates is written under the table rather than
        dropped — losing a sentence to a malformed row would be a far worse
        outcome than an untidy table.
        """
        root_id = cells[0].get("table")
        direct = [block for block in cells if block.get("table") == root_id]
        placed = [block for block in direct if block.get("row") and block.get("cell")]
        stray = [block for block in direct if block not in placed]
        events = []
        nested_seen = set()
        for block in cells:
            if block.get("table") == root_id:
                if block in placed:
                    events.append((block, None))
            elif block.get("parent_table") == root_id and block["table"] not in nested_seen:
                nested_seen.add(block["table"])
                owner = dict(block, row=block["parent_row"], cell=block["parent_cell"],
                             row_span=1, col_span=1)
                events.append((owner, block["table"]))
        positions = [block for block, _ in events]
        if positions:
            # Sized by the grid the cells *cover*, not by the highest
            # coordinate: a cell at row 2 spanning two rows needs a third row
            # to exist before it can be merged into.
            rows = max(int(c["row"]) + int(c.get("row_span", 1)) - 1
                       for c in positions)
            columns = max(int(c["cell"]) + int(c.get("col_span", 1)) - 1
                          for c in positions)
            table = (container.add_table(rows=rows, cols=columns)
                     if container is not None else self.document.add_table(rows=rows, cols=columns))
            if _has_style(self.document, "Table Grid"):
                table.style = "Table Grid"
            if self.options.rtl:
                # Word mirrors the column order for a right-to-left table; a
                # Persian table read left-to-right has its first column last.
                ooxml.set_table_rtl(table)

            written = set()
            for block, nested_id in events:
                top = int(block["row"]) - 1
                left = int(block["cell"]) - 1
                cell = table.cell(top, left)

                # Put a merge back before writing into it: merging afterwards
                # concatenates the paragraphs of every cell involved, so the
                # text would appear once per grid position it covers.
                bottom = top + int(block.get("row_span", 1)) - 1
                right = left + int(block.get("col_span", 1)) - 1
                position = (top, left)
                if position not in written and (bottom, right) != (top, left):
                    cell = cell.merge(table.cell(min(bottom, rows - 1),
                                                 min(right, columns - 1)))
                if position not in written:
                    for paragraph in list(cell.paragraphs):
                        if not paragraph.text and len(paragraph._p) == 0:
                            cell._tc.remove(paragraph._p)
                    written.add(position)
                if nested_id is None:
                    self._cell_event(cell, block)
                else:
                    descendants = {nested_id}
                    for item in cells:
                        if item.get("parent_table") in descendants:
                            descendants.add(item["table"])
                    self._table([item for item in cells if item.get("table") in descendants],
                                container=cell)
                    placeholder = cell.paragraphs[-1] if cell.paragraphs else None
                    if placeholder is not None and not placeholder.text and len(placeholder._p) == 0:
                        cell._tc.remove(placeholder._p)

        for block in stray:
            self._text_block(block)

    def _cell_event(self, cell, block: dict[str, Any]) -> None:
        if block["type"] == "image":
            self._image(block, container=cell)
            return
        paragraph = cell.add_paragraph()
        ooxml.set_paragraph_rtl(paragraph, self.options.rtl)
        if block["type"] == "pagebreak":
            if not block.get("soft") and self.options.page_breaks == "source":
                paragraph.add_run().add_break(WD_BREAK.PAGE)
        else:
            self.write_markup(paragraph, self._text_of(block), block.get("links"))
        self._place_bookmarks(paragraph, block)

    def _text_block(self, block: dict[str, Any]) -> None:
        text = self._text_of(block)
        if not text.strip():
            return
        kind = block["type"]
        if kind == "listitem":
            style = "List Number" if block.get("ordered") else "List Bullet"
            style = style if _has_style(self.document, style) else "Normal"
        else:
            style = _STYLE_BY_TYPE.get(kind, "Normal")
            style = style if _has_style(self.document, style) else "Normal"

        align = None
        if kind == "paragraph" and self.options.justify:
            align = WD_ALIGN_PARAGRAPH.JUSTIFY
        elif kind in ("caption", "verse"):
            align = WD_ALIGN_PARAGRAPH.CENTER

        paragraph = self.paragraph(style, align=align)
        self.write_markup(paragraph, text, block.get("links"))
        self._place_bookmarks(paragraph, block)

    def _text_of(self, block: dict[str, Any]) -> str:
        target = (block.get("target") or "").strip()
        if target:
            return target
        source = (block.get("text") or "").strip()
        if source:
            self.warnings.append(f"block {block['id']} is untranslated")
        return source

    def _image(self, block: dict[str, Any], *, container=None) -> None:
        path = self.assets / block["asset"]
        if not path.exists():
            self.warnings.append(f"missing asset {block['asset']}")
            return

        width_pt, height_pt = self._image_size(block)
        paragraph = (container.add_paragraph() if container is not None else
                     self.paragraph("Normal", align=WD_ALIGN_PARAGRAPH.CENTER))
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        ooxml.set_paragraph_rtl(paragraph, self.options.rtl)
        run = paragraph.add_run()
        try:
            if height_pt:
                run.add_picture(str(path), width=Pt(width_pt), height=Pt(height_pt))
            else:
                run.add_picture(str(path), width=Pt(width_pt))
        except Exception as error:  # unsupported or corrupt image stream
            self.warnings.append(f"could not place {block['asset']}: {error}")
            return

        alt = (block.get("target_alt") or block.get("alt") or "").strip()
        if alt and self.options.captions:
            style = "Caption" if _has_style(self.document, "Caption") else "Normal"
            caption = (container.add_paragraph(style=style) if container is not None else
                       self.paragraph(style, align=WD_ALIGN_PARAGRAPH.CENTER))
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            ooxml.set_paragraph_rtl(caption, self.options.rtl)
            self.write_markup(caption, alt)

    def _image_size(self, block: dict[str, Any]) -> tuple[float, float | None]:
        """Physical size in points, preserving aspect and fitting the text block."""
        available = self.text_width_pt
        width = block.get("width_pt")
        height = block.get("height_pt")

        if not width:
            pixel_width = block.get("pixel_width")
            pixel_height = block.get("pixel_height")
            if pixel_width:
                width = pixel_width * 72.0 / ASSUMED_DPI
                height = (pixel_height * 72.0 / ASSUMED_DPI) if pixel_height else None
            else:
                return available * 0.6, None

        if width > available:
            if height:
                height = height * (available / width)   # aspect ratio is preserved
            width = available
        return width, height

    def _pagebreak(self, block: dict[str, Any]) -> None:
        if self.options.page_breaks != "source" or block.get("soft"):
            return
        paragraph = self.paragraph("Normal")
        paragraph.add_run().add_break(WD_BREAK.PAGE)

    # -- run ---------------------------------------------------------------- #

    def build(self, destination: Path) -> dict[str, Any]:
        self.setup()
        self.collect_toc()
        self.front_matter()
        self.table_of_contents()
        self.body()
        self.footnotes.finalise()

        unused = [n for n in self.notes_by_id if n not in self.used_notes]
        if unused:
            self.warnings.append(
                f"{len(unused)} footnote(s) never referenced: {', '.join(unused[:5])}"
            )

        destination.parent.mkdir(parents=True, exist_ok=True)
        self.document.save(str(destination))
        return {
            "output": str(destination),
            "layout": getattr(self, "layout", {}),
            "headings": len(self.toc_entries),
            "bookmarks": len(self.bookmarks.names),
            "footnotes": len(self.footnotes),
            "images": sum(1 for b in self.book["blocks"] if b["type"] == "image"),
            "warnings": self.warnings[:40],
            "warning_count": len(self.warnings),
        }
